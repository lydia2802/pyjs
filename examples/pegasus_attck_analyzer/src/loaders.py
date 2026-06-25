from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any


def load_json(path: Path) -> Any:
    with path.open("r", encoding="utf-8") as file:
        return json.load(file)


def load_csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        return list(csv.DictReader(file))


def load_xlsx_rows(path: Path) -> dict[str, list[dict[str, Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "Paket 'openpyxl' belum terpasang. Jalankan: pip install -r requirements.txt"
        ) from exc

    workbook = load_workbook(path, data_only=True)
    result: dict[str, list[dict[str, Any]]] = {}

    for sheet in workbook.worksheets:
        headers = [cell.value for cell in sheet[1]]
        rows: list[dict[str, Any]] = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if all(value is None for value in row):
                continue
            item = {
                str(headers[index]) if headers[index] is not None else f"column_{index + 1}": value
                for index, value in enumerate(row)
            }
            rows.append(item)
        result[sheet.title] = rows

    return result


def list_input_files(input_dir: Path) -> list[Path]:
    if not input_dir.exists():
        raise FileNotFoundError(f"Folder input tidak ditemukan: {input_dir}")
    return sorted(path for path in input_dir.iterdir() if path.is_file())

"""
loaders.py
==========

Modul ini bertanggung jawab untuk SEMUA proses pembacaan file mentah
(raw data) dari folder input, yaitu:

1. File JSON       -> ATT&CK Navigator layer (Pegasus Android/iOS) dan changelog.json
2. File CSV        -> known_exploited_vulnerabilities.csv (data KEV dari CISA)
3. File XLSX       -> pegasus_for_android_(s0316).xlsx

Prinsip modul ini:
- Tidak melakukan analisis apa pun (itu tugas analyzers.py).
- Hanya membaca file dari disk dan mengembalikannya dalam bentuk
  struktur data Python standar (dict / list) yang siap diolah.
- Setiap fungsi memvalidasi input dan melempar error yang jelas
  (dalam Bahasa Indonesia) supaya mudah dipahami oleh pengguna.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any


def load_json(path: Path) -> Any:
    """
    Membaca satu file JSON dan mengembalikan isinya sebagai objek Python
    (biasanya dict atau list, tergantung struktur file JSON tersebut).

    Parameter
    ---------
    path : Path
        Lokasi file JSON yang akan dibaca.

    Return
    ------
    Any
        Hasil parsing JSON (umumnya dict).

    Exception
    ---------
    FileNotFoundError
        Jika file tidak ditemukan di lokasi yang diberikan.
    ValueError
        Jika isi file bukan JSON yang valid (rusak/corrupt).
    """
    if not path.exists():
        raise FileNotFoundError(f"File JSON tidak ditemukan: {path}")
    if not path.is_file():
        raise ValueError(f"Path bukan sebuah file: {path}")

    try:
        with path.open("r", encoding="utf-8") as file:
            return json.load(file)
    except json.JSONDecodeError as exc:
        raise ValueError(
            f"Gagal membaca JSON pada file '{path.name}'. "
            f"Pastikan format JSON valid. Detail error: {exc}"
        ) from exc


def load_csv_rows(path: Path) -> list[dict[str, str]]:
    """
    Membaca file CSV dan mengembalikan setiap baris sebagai dict,
    dengan key berupa nama kolom (header baris pertama).

    Menggunakan encoding 'utf-8-sig' agar BOM (Byte Order Mark) yang
    sering muncul pada file CSV hasil export Excel/Windows tidak
    merusak nama kolom pertama.

    Parameter
    ---------
    path : Path
        Lokasi file CSV yang akan dibaca.

    Return
    ------
    list[dict[str, str]]
        Daftar baris data, masing-masing berbentuk dict {nama_kolom: nilai}.

    Exception
    ---------
    FileNotFoundError
        Jika file tidak ditemukan.
    ValueError
        Jika file CSV tidak memiliki header (kosong total).
    """
    if not path.exists():
        raise FileNotFoundError(f"File CSV tidak ditemukan: {path}")
    if not path.is_file():
        raise ValueError(f"Path bukan sebuah file: {path}")

    with path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file)
        rows = list(reader)

        if reader.fieldnames is None:
            raise ValueError(f"File CSV '{path.name}' tidak memiliki baris header.")

    return rows


def load_xlsx_rows(path: Path) -> dict[str, list[dict[str, Any]]]:
    """
    Membaca file Excel (.xlsx) dan mengembalikan isi setiap sheet
    sebagai daftar dict (baris pertama dianggap header kolom).

    Baris yang seluruh selnya kosong (None) akan diabaikan/dilewati.

    Parameter
    ---------
    path : Path
        Lokasi file .xlsx yang akan dibaca.

    Return
    ------
    dict[str, list[dict[str, Any]]]
        Mapping dari nama_sheet -> daftar baris data pada sheet tersebut.

    Exception
    ---------
    FileNotFoundError
        Jika file tidak ditemukan.
    RuntimeError
        Jika library 'openpyxl' belum terpasang di environment Python.
    """
    if not path.exists():
        raise FileNotFoundError(f"File XLSX tidak ditemukan: {path}")
    if not path.is_file():
        raise ValueError(f"Path bukan sebuah file: {path}")

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "Paket 'openpyxl' belum terpasang. "
            "Jalankan: pip install -r requirements.txt"
        ) from exc

    workbook = load_workbook(path, data_only=True)
    result: dict[str, list[dict[str, Any]]] = {}

    for sheet in workbook.worksheets:
        # Baris pertama dianggap sebagai header/nama kolom.
        header_row = sheet[1]
        headers = [cell.value for cell in header_row]

        rows: list[dict[str, Any]] = []
        for raw_row in sheet.iter_rows(min_row=2, values_only=True):
            # Lewati baris yang benar-benar kosong (semua sel None).
            if all(value is None for value in raw_row):
                continue

            row_dict: dict[str, Any] = {}
            for index, value in enumerate(raw_row):
                header_value = headers[index] if index < len(headers) else None
                column_name = str(header_value) if header_value is not None else f"column_{index + 1}"
                row_dict[column_name] = value

            rows.append(row_dict)

        result[sheet.title] = rows

    workbook.close()
    return result


def list_input_files(input_dir: Path) -> list[Path]:
    """
    Mengembalikan daftar semua file (bukan folder) yang ada di dalam
    folder input, terurut berdasarkan nama file.

    Parameter
    ---------
    input_dir : Path
        Folder yang berisi file-file input mentah.

    Return
    ------
    list[Path]
        Daftar path file, terurut alfabetis.

    Exception
    ---------
    FileNotFoundError
        Jika folder input tidak ditemukan.
    NotADirectoryError
        Jika path yang diberikan bukan sebuah folder.
    """
    if not input_dir.exists():
        raise FileNotFoundError(f"Folder input tidak ditemukan: {input_dir}")
    if not input_dir.is_dir():
        raise NotADirectoryError(f"Path input bukan sebuah folder: {input_dir}")

    return sorted(path for path in input_dir.iterdir() if path.is_file())
